"""Import name/phone-only prospects for the existing campaign application flow."""
import io
import re
import secrets

from flask import Blueprint, abort, flash, redirect, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app import db
from app.models import AuditLog, LapsedPolicy, CompanyGroupState
from app.security import permission_required

prospects_bp = Blueprint("prospects", __name__, url_prefix="/prospects")


def prospect_phone(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    from app.routes.communications import _normalise_za_phone
    if re.search(r"[^0-9+()\s.-]", str(value or "")):
        return None
    phone = _normalise_za_phone(str(value or ""))
    # Mobile prefixes, including numbers stored by Excel without their leading zero.
    return phone if phone and phone[3] in "678" else None


@prospects_bp.route("/template.xlsx")
@login_required
@permission_required("recovery.view")
def template():
    wb = Workbook()
    ws = wb.active
    ws.title = "PROSPECT_IMPORT"
    ws.append(["Name", "Cell_Number"])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 25
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="69499D")
    for row in ws.iter_rows(min_row=2, max_row=1001, min_col=2, max_col=2):
        row[0].number_format = "@"
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="Individual_Clients_Template.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@prospects_bp.route("/import", methods=["POST"])
@login_required
@permission_required("recovery.import")
def import_contacts():
    upload = request.files.get("file")
    rows = []
    errors = []
    try:
        if not upload or not (upload.filename or "").lower().endswith(".xlsx"):
            raise ValueError("Choose an .xlsx file using the individual clients template.")
        if len(upload.read()) > 5 * 1024 * 1024:
            raise ValueError("The import file must be smaller than 5 MB.")
        upload.seek(0)
        wb = load_workbook(upload, read_only=True, data_only=True)
        try:
            ws = wb.active
            iterator = ws.iter_rows(values_only=True)
            headers = [re.sub(r"[^a-z0-9]", "", str(v or "").lower()) for v in next(iterator, ())]
            if headers != ["name", "cellnumber"]:
                raise ValueError("The first row must contain exactly Name and Cell_Number.")
            for number, row in enumerate(iterator, 2):
                if number > 10001:
                    raise ValueError("Import at most 10,000 contacts per file.")
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                name = str(row[0] or "").strip()
                phone = prospect_phone(row[1] if len(row) > 1 else None)
                parts = name.rsplit(None, 1)
                first = parts[0] if parts else ""
                surname = parts[1] if len(parts) > 1 else ""
                if not first or len(first) > 50 or len(surname) > 120 or not phone:
                    errors.append(f"Row {number}: enter a name (first names up to 50 characters, surname up to 120) and a valid SA mobile number.")
                else:
                    rows.append((first, surname, phone))
        finally:
            wb.close()
        if errors:
            raise ValueError("No contacts imported. " + " ".join(errors[:10]))
        if not rows:
            raise ValueError("The file contains no contacts.")
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else "The workbook could not be read. Use the blank individual clients template."
        flash(message, "danger")
        return redirect(url_for("recovery.queue"))

    from app.routes.communications import _available_companies
    from app.services.company_groups import get_or_create_company
    company_id = request.form.get("company_id", type=int)
    company = db.session.get(CompanyGroupState, company_id) if company_id else None
    if company_id and (not company or company.id not in {c.id for c in _available_companies()}):
        abort(403)
    if company and company.status != "Active":
        flash("Choose an active company for this import.", "danger")
        return redirect(url_for("recovery.queue"))
    created = skipped = 0
    seen = set()
    from app.services.communication_service import normalize_phone
    try:
        if not company:
            company = get_or_create_company("Individual Clients", current_user.branch or "")
        if company.status != "Active":
            db.session.rollback()
            flash("Choose an active company for this import.", "danger")
            return redirect(url_for("recovery.queue"))
        from app.routes.communications import _company_leads
        existing = {normalize_phone(p.cell_number) for p in _company_leads(
            LapsedPolicy.query, company).all()}
        for first, surname, phone in rows:
            key = normalize_phone(phone)
            if key in seen or key in existing:
                skipped += 1
                continue
            seen.add(key)
            lead = LapsedPolicy(lead_type="prospect", member_id="PROSPECT-" + secrets.token_hex(12),
                                initials=first, surname=surname, cell_number=phone,
                                branch=company.branch, company_id=company.id, company_name=company.company_name,
                                franchise=company.company_name, assigned_agent_id=current_user.id,
                                recovery_status="New", next_action_date=None,
                                comments="Imported from the name and cell number template.")
            db.session.add(lead)
            created += 1
        db.session.add(AuditLog(user_id=current_user.id, action="PROSPECT_IMPORT",
            entity_type="CompanyGroupState", entity_id=str(company.id),
            details=f"Imported {created} name/phone contacts; skipped {skipped} duplicates."))
        db.session.commit()
    except Exception:
        db.session.rollback()
        flash("No contacts imported. The database could not save this import; please retry.", "danger")
        return redirect(url_for("recovery.queue"))
    flash(f"Imported {created} individual clients; skipped {skipped} duplicate contacts. Create a campaign and select Individual clients only.", "success")
    return redirect(url_for("recovery.queue"))


@prospects_bp.app_context_processor
def import_companies():
    if current_user.is_authenticated and request.endpoint == "recovery.queue":
        from app.routes.communications import _available_companies
        return {"prospect_companies": _available_companies()}
    return {"prospect_companies": []}
