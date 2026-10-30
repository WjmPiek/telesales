import io
import json
import unittest
from datetime import date, datetime
from unittest.mock import patch
from openpyxl import Workbook
import test_analytics_security as fixtures

from app import db
from app.models import (LapsedPolicy, ClientApplication, ApplicationJourney, PolicyProduct,
                        ContactSuppression, CampaignRecipient, CommunicationCampaign)
from app.services.communication_service import contact_hash, normalize_phone, preference_for
from app.services.client_directory import filter_policy_query, scoped_policies


class LapsReportTests(unittest.TestCase):
    setUp = fixtures.AnalyticsSecurityTests.setUp
    tearDown = fixtures.AnalyticsSecurityTests.tearDown
    login = fixtures.AnalyticsSecurityTests.login

    @staticmethod
    def workbook(rows=None):
        workbook = Workbook()
        workbook.active.append(['Policy #', 'Receipt #', 'Initial. Surname', 'Address', 'Cell#', 'Phone#', 'Last Receipt', 'MPIA', 'Comment'])
        for row in rows if rows is not None else [
            ['TEST-1', '123', 'JP VAN DER TEST', '1 Fictional Road', 821234567, '0', datetime(2023, 6, 22), 1, 'Source note'],
            ['TEST-2', 'NONE YET', 'A EXAMPLE', '2 Fictional Road', '0', '', datetime(1900, 1, 1), 0, 'No receipt'],
        ]:
            workbook.active.append(row)
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def import_report(self, filename='Northcliff Laps Report.xlsx', rows=None):
        self.login('admin')
        return self.client.post('/recovery/import', data={'file': (self.workbook(rows), filename)})

    def lead(self, policy='TEST-1', franchise='Northcliff'):
        return LapsedPolicy.query.filter_by(policy_number=policy, franchise=franchise).one()

    def test_exact_report_headers_preserve_details_and_allow_no_id(self):
        self.assertEqual(self.import_report().status_code, 302)
        lead = self.lead()
        self.assertEqual((lead.initials, lead.surname, lead.address), ('JP', 'VAN DER TEST', '1 Fictional Road'))
        self.assertEqual(normalize_phone(lead.cell_number), '27821234567')
        self.assertEqual(lead.last_date_paid, date(2023, 6, 22))
        self.assertEqual(lead.receipt_number, '123')
        self.assertEqual(lead.source_filename, 'Northcliff Laps Report.xlsx')
        self.assertIsNone(lead.id_number)
        self.assertEqual(lead.recovery_status, 'Imported')
        original = json.loads(lead.source_row_json)
        self.assertEqual(original['MPIA'], 1)
        self.assertEqual(original['Initial. Surname'], 'JP VAN DER TEST')
        from app.routes.recovery import _policy_missing_fields
        self.assertEqual(_policy_missing_fields(lead), [])
        missing = self.lead('TEST-2')
        self.assertEqual(missing.recovery_status, 'Suspense')
        self.assertIsNone(missing.last_date_paid)
        self.assertIn('1900-01-01', missing.source_row_json)

    def test_named_franchise_uses_existing_company_instead_of_admin_branch(self):
        from app.models import CompanyGroupState
        company = CompanyGroupState(company_name='MARTINS FUNERALS NORTHCLIFF', branch='NORTHCLIFF', status='Active')
        db.session.add(company); db.session.commit()
        company_id = company.id
        self.import_report()
        lead = self.lead()
        self.assertEqual((lead.company_id, lead.branch), (company_id, 'NORTHCLIFF'))
        campaign = CommunicationCampaign(name='Franchise campaign',message_body='Test',company_id=company_id,branch='NORTHCLIFF',created_by_id=self.user_ids['admin'])
        db.session.add(campaign); db.session.commit()
        response = self.client.get(f'/communications/{campaign.id}?date_basis=imported&date_from={date.today()}&date_to={date.today()}')
        self.assertIn(b'VAN DER TEST',response.data)
        self.assertIn(b'Northcliff',response.data)

    def test_group_repair_is_explicit_authorized_and_keeps_client_ids(self):
        from app.models import CompanyGroupState
        self.import_report()
        lead_ids = [self.lead().id, self.lead('TEST-2').id]
        company = CompanyGroupState(company_name='MARTINS FUNERALS NORTHCLIFF',branch='NORTHCLIFF',status='Active')
        db.session.add(company); db.session.commit()
        self.assertEqual(self.client.post('/settings/report-import-groups').status_code,400)
        page=self.client.get('/settings/report-import-groups')
        self.assertIn(b'Correct imported franchise groups',page.data)
        with self.client.session_transaction() as session:
            csrf=session['report_group_csrf']
        self.client.post('/settings/report-import-groups',data={'csrf':csrf})
        self.assertEqual([self.lead().id,self.lead('TEST-2').id],lead_ids)
        self.assertEqual(self.lead().company_id,company.id)
        self.assertEqual(self.lead().branch,'NORTHCLIFF')
        self.import_report()
        self.assertEqual(LapsedPolicy.query.filter_by(source_type='laps_report').count(),2)
        self.login('manager')
        self.assertEqual(self.client.get('/settings/report-import-groups').status_code,403)

    def test_report_mapping_rejects_ambiguous_companies_and_cross_branch_staff(self):
        from app.models import CompanyGroupState
        from app.services.laps_report import report_company, import_report
        from openpyxl import load_workbook
        db.session.add(CompanyGroupState(company_name='MARTINS FUNERALS NORTHCLIFF',branch='NORTHCLIFF',status='Active'))
        db.session.commit()
        wb=load_workbook(self.workbook())
        with self.assertRaisesRegex(ValueError,'another branch'):
            import_report(wb,'Northcliff Laps Report.xlsx',self.users['manager'])
        wb.close()
        db.session.add(CompanyGroupState(company_name='Northcliff',branch='NORTHCLIFF',status='Active'))
        db.session.commit()
        with self.assertRaisesRegex(ValueError,'More than one company'):
            report_company('Northcliff','Brokers')

    def test_group_repair_skips_clients_already_in_a_campaign(self):
        from app.models import CompanyGroupState
        from app.services.laps_report import report_group_repairs
        self.import_report()
        lead_id=self.lead().id
        company=CompanyGroupState(company_name='MARTINS FUNERALS NORTHCLIFF',branch='NORTHCLIFF',status='Active')
        campaign=CommunicationCampaign(name='Existing',message_body='Test',created_by_id=self.user_ids['admin'])
        db.session.add_all([company,campaign]); db.session.flush()
        db.session.add(CampaignRecipient(campaign_id=campaign.id,lapsed_policy_id=lead_id,secure_token='repair-test'))
        db.session.commit()
        self.assertEqual(report_group_repairs(),[])

    def test_reimport_updates_in_place_without_resetting_progress_or_opt_out(self):
        self.import_report()
        lead = self.lead()
        lead.recovery_status = 'Application Started'
        lead_id = lead.id
        preference_for(lead).opted_out_all = True
        db.session.commit()
        self.import_report()
        lead = self.lead()
        self.assertEqual(lead.id, lead_id)
        self.assertEqual(LapsedPolicy.query.filter_by(source_type='laps_report').count(), 2)
        self.assertEqual(lead.recovery_status, 'Opted Out')
        self.assertTrue(preference_for(lead).opted_out_all)

    def test_same_policy_number_in_different_franchises_stays_separate(self):
        self.import_report()
        self.import_report('Southridge Laps Report.xlsx')
        self.assertNotEqual(self.lead().id, self.lead(franchise='Southridge').id)
        self.assertEqual(LapsedPolicy.query.filter_by(source_type='laps_report').count(), 4)

    def test_historical_suppression_is_preserved_on_import(self):
        db.session.add(ContactSuppression(phone_hash=contact_hash('27821234567')))
        db.session.commit()
        self.import_report()
        lead = self.lead()
        self.assertEqual(lead.recovery_status, 'Opted Out')
        self.assertFalse(preference_for(lead).whatsapp_allowed)

    def test_duplicate_or_malformed_report_does_not_partially_import(self):
        rows = [['TEST-1', '1', 'A EXAMPLE', '', '0821234567', '', date(2023, 6, 1), 1, '']]
        self.import_report(rows=rows + rows)
        self.assertEqual(LapsedPolicy.query.filter_by(source_type='laps_report').count(), 0)
        rows.append(['TEST-2', '2', 'B EXAMPLE', '', '0821234568', '', 'not a date', 1, ''])
        self.import_report(rows=rows)
        self.assertEqual(LapsedPolicy.query.filter_by(source_type='laps_report').count(), 0)

    def test_directory_filters_franchise_receipt_year_month_and_range(self):
        self.import_report()
        self.import_report('Southridge Laps Report.xlsx', [['S-1', '2', 'B EXAMPLE', '', '0821234568', '', date(2024, 7, 1), 1, '']])
        from flask import g
        self.login('admin')
        with self.app.test_request_context():
            g.pop('_login_user', None)
            from flask_login import login_user
            from app.models import User
            login_user(db.session.get(User, self.user_ids['admin']))
            base = scoped_policies()
            matches = filter_policy_query(base, {'franchise':'Northcliff','year':'2023','month':'6'}).all()
            self.assertEqual([p.policy_number for p in matches], ['TEST-1'])
            matches = filter_policy_query(base, {'year':'2024','date_from':'2024-07-01','date_to':'2024-07-01'}).all()
            self.assertEqual([p.policy_number for p in matches], ['S-1'])
        page = self.client.get('/client-files/?franchise=Northcliff&year=2023&month=6')
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'TEST-1', page.data)
        self.assertNotIn(b'S-1', page.data)
        self.assertIn(b'All franchises combined', page.data)
        self.assertNotIn(b'Download blank', self.client.get('/recovery/').data)

    def test_process_status_is_based_on_active_new_policy_and_keeps_removal(self):
        self.import_report()
        lead_id = self.lead().id
        application = ClientApplication(application_ref='JOIN-TEST', lapsed_policy_id=lead_id, branch='A',
                                        agent_id=self.user_ids['admin'], application_type='New Policy - Self Service', status='Signed')
        db.session.add(application)
        db.session.commit()
        self.assertIn(b'TEST-1', self.client.get('/client-files/?process=in_progress').data)
        self.assertNotIn(b'TEST-1', self.client.get('/client-files/?process=completed').data)
        application = ClientApplication.query.filter_by(application_ref='JOIN-TEST').one()
        application.status = 'Active'
        application.inception_date = date(2026, 10, 1)
        db.session.commit()
        self.assertIn(b'TEST-1', self.client.get('/client-files/?process=completed&date_basis=joined&year=2026&month=10').data)
        preference_for(self.lead()).opted_out_all = True
        db.session.commit()
        self.assertIn(b'TEST-1', self.client.get('/client-files/?process=removed').data)

    def test_directory_and_filter_choices_keep_branch_access(self):
        self.import_report()
        self.login('other')
        page = self.client.get('/client-files/?franchise=Northcliff')
        self.assertEqual(page.status_code, 200)
        self.assertNotIn(b'TEST-1', page.data)
        self.assertNotIn(b'<option value="Northcliff"', page.data)
        self.login('agent')
        self.assertNotIn(b'TEST-1', self.client.get('/client-files/').data)

    def test_import_dates_use_south_african_month_boundary(self):
        self.import_report()
        self.lead().imported_at = datetime(2025, 12, 31, 22, 30)
        db.session.commit()
        response = self.client.get('/client-files/?date_basis=imported&year=2026&month=1')
        self.assertIn(b'TEST-1', response.data)
        self.assertNotIn(b'TEST-1', self.client.get('/client-files/?date_basis=imported&year=2025&month=12').data)

    def test_campaign_filtered_group_uses_same_franchise_and_dates(self):
        self.import_report()
        self.import_report('Southridge Laps Report.xlsx', [['S-1', '2', 'B EXAMPLE', '', '0821234568', '', date(2024, 7, 1), 1, '']])
        campaign = CommunicationCampaign(name='Filtered group', message_body='Test only', created_by_id=self.user_ids['admin'],
                                         audience_type='group', branch='A')
        db.session.add(campaign)
        db.session.commit()
        campaign_id = campaign.id
        page = self.client.get(f'/communications/{campaign_id}?franchise=Northcliff&year=2023&month=6')
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'name="franchise"', page.data)
        response = self.client.post(f'/communications/{campaign_id}/add-filtered-group', data={'franchise':'Northcliff','year':'2023','month':'6'})
        self.assertEqual(response.status_code, 302)
        recipients = CampaignRecipient.query.filter_by(campaign_id=campaign_id).all()
        self.assertEqual([r.policy.policy_number for r in recipients], ['TEST-1'])

    def test_submitted_form_merges_into_import_and_retains_original_receipt(self):
        from app.services.online_application import save_questionnaire
        from app.services.cdd_service import FIELDS as CDD_FIELDS
        self.import_report()
        lead = self.lead()
        source_json = lead.source_row_json
        lead_id = lead.id
        product = PolicyProduct(product_name='Test family', plan_name='Test', monthly_premium=100, cover_amount=10000)
        application = ClientApplication(application_ref='MERGE-TEST', lapsed_policy_id=lead.id, company_id=lead.company_id,
                                        branch=lead.branch, agent_id=self.user_ids['admin'], product=product,
                                        first_names=lead.initials, surname=lead.surname, id_number='8001015009087',
                                        application_type='New Policy - Self Service', status='Draft',
                                        cell_number=lead.cell_number, address=lead.address, residential_address=lead.address,
                                        date_of_birth='01/01/1980', form_template='single_family', total_members=1,
                                        monthly_premium=100, cover_amount=10000)
        db.session.add(application)
        db.session.flush()
        application_id = application.id
        db.session.add(ApplicationJourney(application_id=application.id))
        db.session.commit()
        data = {'first_names':'James Paul', 'surname':'VAN DER TEST', 'cell_number':'0827654321',
                'email':'client@example.test','residential_address':'99 Updated Road','residential_postal_code':'1234',
                'postal_address':'PO Box 1','postal_code':'1234','beneficiary_full_names':'Example Beneficiary',
                'beneficiary_relationship':'Spouse','beneficiary_date_of_birth':'1985-01-01',
                'payment_method':'Cash','marketing_choice':'yes'}
        for key,label,options in CDD_FIELDS:
            data['cdd_'+key] = options.split('|')[0] if options else 'Test answer'
        application = db.session.get(ClientApplication, application_id)
        count_before = LapsedPolicy.query.count()
        save_questionnaire(application, data)
        db.session.commit()
        self.assertEqual(LapsedPolicy.query.count(), count_before)
        lead = db.session.get(LapsedPolicy, lead_id)
        self.assertEqual((lead.initials, lead.surname, lead.address, lead.cell_number, lead.email_address),
                         ('James Paul','VAN DER TEST','99 Updated Road','0827654321','client@example.test'))
        self.assertEqual(lead.id_number, '8001015009087')
        self.assertEqual(lead.last_date_paid, date(2023, 6, 22))
        self.assertEqual(lead.source_row_json, source_json)
        self.assertEqual(application.beneficiary_full_names, 'Example Beneficiary')
        self.assertEqual(application.postal_address, 'PO Box 1')
        self.import_report()
        lead = db.session.get(LapsedPolicy, lead_id)
        self.assertEqual(lead.address, '99 Updated Road')
        self.assertEqual(lead.initials, 'James Paul')
        page = self.client.get(f'/client-files/?application_id={application_id}')
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'Original imported report details', page.data)
        self.assertIn(b'1 Fictional Road', page.data)
        self.assertIn(b'99 Updated Road', page.data)

    def test_application_file_follows_explicit_import_link_before_id_merge(self):
        self.import_report()
        lead = self.lead()
        a = ClientApplication(application_ref='UNMERGED', lapsed_policy_id=lead.id, id_number='8001015009087',
                              branch='A', agent_id=self.user_ids['admin'])
        db.session.add(a)
        db.session.commit()
        page = self.client.get(f'/client-files/?application_id={a.id}')
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'Northcliff Laps Report.xlsx', page.data)

    def test_invalid_filters_are_rejected(self):
        self.login('admin')
        for query in ('month=13','year=nope','date_from=2026-12-01&date_to=2026-01-01','process=unknown'):
            self.assertEqual(self.client.get('/client-files/?'+query).status_code, 400)

    def test_report_columns_upgrade_existing_table_without_losing_rows(self):
        from app.services.laps_report import ensure_report_columns
        from sqlalchemy import create_engine, text, inspect
        # Exercise the same application-bound migration against an old schema.
        with db.engine.begin() as conn:
            conn.execute(text('DROP TABLE lapsed_policies'))
            conn.execute(text('CREATE TABLE lapsed_policies (id INTEGER PRIMARY KEY, policy_number VARCHAR(80))'))
            conn.execute(text("INSERT INTO lapsed_policies VALUES (1,'OLD-REPORT')"))
        ensure_report_columns()
        ensure_report_columns()
        self.assertTrue({'source_type','source_filename','receipt_number','source_row_json'} <=
                        {column['name'] for column in inspect(db.engine).get_columns('lapsed_policies')})
        with db.engine.connect() as conn:
            self.assertEqual(conn.execute(text('SELECT policy_number FROM lapsed_policies')).scalar(), 'OLD-REPORT')

    def test_trailing_formatted_blank_columns_do_not_break_report_import(self):
        from openpyxl.styles import PatternFill
        from openpyxl import load_workbook
        workbook=load_workbook(self.workbook())
        workbook.active['K1'].fill=PatternFill('solid',fgColor='FFFFFF')
        output=io.BytesIO();workbook.save(output);output.seek(0)
        self.login('admin')
        response=self.client.post('/recovery/import',data={'file':(output,'Northcliff Laps Report.xlsx')})
        self.assertEqual(response.status_code,302)
        self.assertEqual(LapsedPolicy.query.filter_by(source_type='laps_report').count(),2)

    def test_unknown_and_empty_workbooks_are_not_reported_as_successful_imports(self):
        self.login('admin')
        for headers in [['Unrelated column'],['Policy_Number','Surname','Cell_Number']]:
            workbook=Workbook();workbook.active.append(headers)
            output=io.BytesIO();workbook.save(output);output.seek(0)
            response=self.client.post('/recovery/import',data={'file':(output,'unrecognised.xlsx')})
            self.assertEqual(response.status_code,302)
            self.assertEqual(LapsedPolicy.query.count(),2)
            with self.client.session_transaction() as session:
                self.assertEqual(session['_flashes'][-1][0],'danger')
