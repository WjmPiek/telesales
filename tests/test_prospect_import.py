import io
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import test_application_flow as fixtures
from app import db
from app.models import (AgentNotification, CampaignRecipient, ClientApplication, CommunicationCampaign,
                        CompanyGroupState, LapsedPolicy, User)
from app.services.cdd_service import FIELDS as CDD_FIELDS
from app.services.communication_service import preference_for
from app.services.whatsapp_service import SendResult


class ProspectImportTests(unittest.TestCase):
    setUp = fixtures.ApplicationFlowTests.setUp
    tearDown = fixtures.ApplicationFlowTests.tearDown

    def workbook(self, rows, headers=("Name", "Cell_Number")):
        wb = Workbook()
        wb.active.append(headers)
        for row in rows:
            wb.active.append(row)
        data = io.BytesIO()
        wb.save(data)
        data.seek(0)
        return data

    def import_rows(self, rows, company=None, headers=("Name", "Cell_Number")):
        data = {"file": (self.workbook(rows, headers), "contacts.xlsx")}
        if company:
            data["company_id"] = str(company.id)
        return self.client.post("/prospects/import", data=data, content_type="multipart/form-data")

    def campaign(self, lead):
        campaign = CommunicationCampaign(name="Prospects", message_body="Hello {{1}}", created_by_id=self.user_id,
            company_id=lead.company_id, branch=lead.branch, send_email=False, audience_type="group",
            whatsapp_template_name="test_approved", image_url="https://example.test/image.png", template_status="Approved",
            template_buttons_json='[{"type":"URL","text":"APPLY NOW","url":"https://example.test/join/{{1}}"},{"type":"QUICK_REPLY","text":"CALL ME BACK"},{"type":"QUICK_REPLY","text":"DELETE MY NUMBER"}]')
        recipient = CampaignRecipient(campaign=campaign, policy=lead, secure_token="prospect-test-token")
        db.session.add_all([campaign, recipient])
        db.session.commit()
        return campaign, recipient

    def test_template_and_import_without_policy_details(self):
        response = self.client.get("/prospects/template.xlsx")
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.data))
        self.assertEqual([c.value for c in wb.active[1]], ["Name", "Cell_Number"])
        self.assertEqual(wb.active["B2"].number_format, "@")
        self.assertEqual(self.import_rows([("Jane Example", "0676200748")]).status_code, 302)
        lead = LapsedPolicy.query.one()
        self.assertEqual((lead.lead_type, lead.cell_number, lead.recovery_status), ("prospect", "+27676200748", "New"))
        self.assertIsNone(lead.id_number)
        self.assertIsNone(lead.email_address)
        self.assertIsNone(lead.policy_number)
        self.assertEqual(CampaignRecipient.query.count(), 0)
        self.assertEqual(ClientApplication.query.filter_by(lapsed_policy_id=lead.id).count(), 0)
        self.client.post("/recovery/suspense/rebuild")
        self.assertEqual(lead.recovery_status, "New")

    def test_duplicate_formats_are_skipped_without_changing_existing_details(self):
        self.import_rows([("Jane Example", "0676200748"), ("Other Name", "+27 67 620 0748"), ("Numeric Client", 826543210)])
        self.import_rows([("Changed Name", "0027676200748")])
        self.assertEqual(LapsedPolicy.query.count(), 2)
        self.assertEqual(LapsedPolicy.query.filter_by(cell_number="+27676200748").one().initials, "Jane")
        self.assertEqual(LapsedPolicy.query.filter_by(cell_number="+27826543210").one().initials, "Numeric")

    def test_invalid_rows_fail_atomically_and_wrong_template_is_rejected(self):
        self.import_rows([("Valid Client", "0676200748"), ("Broken Client", "123")])
        self.assertEqual(LapsedPolicy.query.count(), 0)
        self.import_rows([("A", "0676200748")], headers=("Policy_Number", "Cell_Number"))
        self.assertEqual(LapsedPolicy.query.count(), 0)

    def test_agent_cannot_import_without_import_permission(self):
        db.session.get(User, self.user_id).role.name = "Agent"
        db.session.commit()
        self.assertEqual(self.import_rows([("Jane Example", "0676200748")]).status_code, 403)
        self.assertEqual(LapsedPolicy.query.count(), 0)

    def test_selected_company_and_campaign_filter_keep_existing_policies_separate(self):
        company = CompanyGroupState(company_name="Business A", branch="A")
        db.session.add(company); db.session.commit()
        company_id = company.id
        self.import_rows([("Jane Example", "0676200748")], company)
        lead = LapsedPolicy.query.one()
        policy = LapsedPolicy(company_id=company_id, branch="A", initials="Existing", surname="Policy",
                              cell_number="0821234567", id_number="8001015009087")
        db.session.add(policy); db.session.commit()
        campaign, recipient = self.campaign(lead)
        page = self.client.get(f"/communications/{campaign.id}?lead_type=prospect")
        self.assertEqual(page.status_code, 200)
        self.assertNotIn(b"Existing Policy", page.data)
        self.client.post(f"/communications/{campaign.id}/add-filtered-group", data={"lead_type": "prospect"})
        self.assertEqual(CampaignRecipient.query.filter_by(campaign_id=campaign.id).count(), 1)

    def test_prospect_campaign_uses_existing_apply_now_template(self):
        self.import_rows([("Jane Example", "0676200748")])
        campaign, recipient = self.campaign(LapsedPolicy.query.one())
        from app.routes.communications import _send_to_recipient
        with patch("app.routes.communications.send_whatsapp_template_image", return_value=SendResult(True, message_id="test-message")) as send:
            ok, error = _send_to_recipient(campaign, recipient, "whatsapp")
        self.assertTrue(ok, error)
        self.assertEqual(send.call_args.kwargs["join_token"], "prospect-test-token")
        self.assertEqual(recipient.whatsapp_status, "Sent")

    def submit_application(self, marketing="yes"):
        self.import_rows([("Jane Example", "0676200748")])
        lead = LapsedPolicy.query.one()
        campaign, recipient = self.campaign(lead)
        public = self.app.test_client()
        self.assertEqual(public.get("/join/prospect-test-token").status_code, 200)
        response = public.post("/join/prospect-test-token", data={"id_number": "8001015009087", "total_members": "4", "cover_amount": "10000"})
        self.assertIn("/products", response.location)
        product_id = fixtures.db.session.get(ClientApplication, self.record_id).product_id
        response = public.get(f"/join/prospect-test-token/application/{product_id}")
        application = ClientApplication.query.filter_by(lapsed_policy_id=lead.id).one()
        self.assertEqual(public.get(response.location).status_code, 200)
        with public.session_transaction() as session:
            nonce = session[f"questionnaire_nonce_{application.id}"]
        data = {"nonce": nonce, "action": "save", "first_names": "Jane", "surname": "Example",
                "cell_number": "0676200748", "email": "jane@example.test", "residential_address": "1 Example Road",
                "residential_postal_code": "1234", "beneficiary_full_names": "Example Beneficiary",
                "beneficiary_relationship": "Spouse", "beneficiary_date_of_birth": "1985-01-01",
                "payment_method": "Cash", "marketing_choice": marketing}
        for key, label, options in CDD_FIELDS:
            data["cdd_" + key] = options.split("|")[0] if options else "Fictional answer"
        path = "/online-application/" + application.sign_token
        response = public.post(path, data=data)
        self.assertEqual(response.status_code, 302, response.data[:400])
        return LapsedPolicy.query.one(), ClientApplication.query.filter_by(source_campaign_recipient_id=recipient.id).one(), public, path, data

    def test_client_submission_saves_details_and_continues_existing_signing_flow(self):
        lead, application, public, path, data = self.submit_application()
        lead_id = lead.id
        self.assertTrue(application.whatsapp_journey.ready)
        self.assertEqual((lead.id_number, lead.email_address, lead.address), ("8001015009087", "jane@example.test", "1 Example Road"))
        self.assertEqual(lead.recovery_status, "Application Started")
        self.assertEqual(AgentNotification.query.filter_by(notification_type="application_submitted").count(), 1)
        self.assertEqual(public.get(f"/sign/{application.sign_token}/review/application").status_code, 200)
        public.post(path, data=data)
        self.assertEqual(AgentNotification.query.filter_by(notification_type="application_submitted").count(), 1)
        self.assertEqual(ClientApplication.query.filter_by(lapsed_policy_id=lead_id).count(), 1)

    def test_marketing_opt_out_is_preserved_when_application_is_saved(self):
        lead, application, public, path, data = self.submit_application(marketing="no")
        self.assertTrue(application.whatsapp_journey.ready)
        self.assertEqual(lead.recovery_status, "Opted Out")
        self.assertTrue(preference_for(lead).opted_out_all)


if __name__ == "__main__":
    unittest.main()
